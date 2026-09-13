from pathlib import Path
from datetime import datetime
import ast
import io
import pandas as pd
import pytest
from gp_core import *
from gp_store import Store,utcnow
from gp_reports import excel_bytes
from gp_ocr import _run_ocr_worker,_ocr_scan_lock
from invoice_ocr_worker import parse_codes,parse_quantity,parse_summary,parse_header_metadata
from PIL import Image


def test_stock_reader_preserves_string_codes():
    df=pd.DataFrame({COL_CODE:["010716"],COL_NAME:["Camera"],COL_QTY:[3]})
    b=io.BytesIO();df.to_excel(b,index=False)
    parsed=read_stock_report(b.getvalue())
    assert parsed.iloc[0][COL_CODE]=="010716"


def test_analysis_and_zero_previous_rate():
    stock=ensure_unique_stock_keys(pd.DataFrame([{COL_CODE:"A",COL_NAME:"Test",COL_QTY:10}]))
    history=pd.DataFrame([{COL_CODE:"A",COL_NAME:"Test",COL_MATCH:"test",COL_DATE:pd.Timestamp("2026-09-01"),
                          COL_REF:"sale",COL_IN:0,COL_OUT:10}])
    output=build_inventory_analysis(stock,history,30,15,90)
    assert len(output)==1


def test_warehouse_master_rejects_missing_code():
    with pytest.raises(ValueError,match="رمز"):
        ensure_unique_stock_keys(pd.DataFrame([{COL_CODE:"",COL_NAME:"Camera",COL_QTY:1}]))


def test_report_export_blocks_formula_injection():
    from openpyxl import load_workbook
    data=excel_bytes({"Test":pd.DataFrame({"name":["=HYPERLINK(\"https://example.com\")"],"qty":[2]})})
    book=load_workbook(io.BytesIO(data));cell=book.active["A2"]
    assert cell.data_type!="f" and cell.value.startswith("'")
    assert book.active.freeze_panes=="A2"


def test_history_import_no_stock_deduction():
    s=Store.for_tests();pw="Test-Password-2026";s.initialize(password=pw);t=s.login("admin",pw)
    stock=ensure_unique_stock_keys(pd.DataFrame([{COL_CODE:"010716",COL_NAME:"Camera",COL_QTY:10}]))
    s.replace_stock(t,pw,stock,"base",s.state(t)["revision"])
    history=pd.DataFrame([{COL_CODE:"010716",COL_NAME:"Camera",COL_MATCH:"CODE:010716",COL_DATE:pd.Timestamp("2026-01-01"),
                          COL_REF:"R1",COL_IN:100,COL_OUT:90,COL_BAL:10}])
    s.import_history(t,pw,history,"a"*64,utcnow(),"history")
    assert s.stock(t).iloc[0][COL_QTY]==10
    assert s.stock(t).iloc[0][COL_CODE]=="010716"
    assert len(s.movement_history(t))==1


def test_import_float_tail_normalizes_to_four_places():
    s=Store.for_tests();pw="Test-Password-2026";s.initialize(password=pw);t=s.login("admin",pw)
    stock=ensure_unique_stock_keys(pd.DataFrame([{COL_CODE:"010716",COL_NAME:"Camera",COL_QTY:3.0000000000000004}]))
    s.replace_stock(t,pw,stock,"base-float",s.state(t)["revision"])
    assert s.stock(t).iloc[0][COL_QTY]==3


def test_codes_no_fuzzy_padding():
    boxes=[{"text":"010716","score":.9,"x":900,"y":400},
           {"text":"abc","score":.9,"x":900,"y":500}]
    rows=parse_codes(boxes,1000,1200)
    assert len(rows)==1 and rows[0]["code"]=="010716"


def test_quantity_ambiguity_blank():
    boxes=[{"text":"3.00","score":.9,"x":280,"y":400},{"text":"4.00","score":.9,"x":281,"y":401}]
    assert parse_quantity(boxes,400,1000,1200) is None


def test_quantity_parser():
    boxes=[{"text":"3.00","score":.9,"x":280,"y":400}]
    assert parse_quantity(boxes,400,1000,1200)==3


def test_reference_parser_prefers_template_position_over_total_score():
    # Real Golden Palace layout: invoice number is above the lower printed total.
    # The total may have a higher OCR confidence and must still not become "900".
    boxes=[{"text":"8042","score":.72,"x":110,"y":620,"region":"summary"},
           {"text":"9.00","score":.98,"x":110,"y":690,"region":"summary"}]
    reference,total=parse_summary(boxes,1000,1200)
    assert reference=="8042" and total==9.0



def test_delete_stock_report_rebuilds_then_clears_current_stock():
    s=Store.for_tests();pw="Test-Password-2026";s.initialize(password=pw);t=s.login("admin",pw)
    first=ensure_unique_stock_keys(pd.DataFrame([{COL_CODE:"010716",COL_NAME:"Camera",COL_QTY:10}]))
    first_op=s.replace_stock(t,pw,first,"base-delete-1",s.state(t)["revision"],source_name="first.xlsx")
    second=ensure_unique_stock_keys(pd.DataFrame([{COL_CODE:"010716",COL_NAME:"Camera",COL_QTY:25}]))
    second_op=s.replace_stock(t,pw,second,"base-delete-2",s.state(t)["revision"],source_name="second.xlsx")
    assert s.stock(t).iloc[0][COL_QTY]==25
    s.delete_stock_report(t,pw,second_op)
    assert s.stock(t).iloc[0][COL_QTY]==10
    s.delete_stock_report(t,pw,first_op)
    assert s.stock(t).empty



def test_stock_reader_hides_orphan_balance_without_warehouse_report():
    from sqlalchemy import delete
    s=Store.for_tests();pw="Test-Password-2026";s.initialize(password=pw);t=s.login("admin",pw)
    stock=ensure_unique_stock_keys(pd.DataFrame([{COL_CODE:"010716",COL_NAME:"Camera",COL_QTY:10}]))
    s.replace_stock(t,pw,stock,"base-orphan-read",s.state(t)["revision"],source_name="orphan.xlsx")
    with s.engine.begin() as c:c.execute(delete(s.tables["baseline_snapshots"]))
    assert s.stock(t).empty


def test_initialize_repairs_orphan_stock_left_by_old_delete_behavior():
    from sqlalchemy import delete,select,func
    s=Store.for_tests();pw="Test-Password-2026";s.initialize(password=pw);t=s.login("admin",pw)
    stock=ensure_unique_stock_keys(pd.DataFrame([{COL_CODE:"010716",COL_NAME:"Camera",COL_QTY:10}]))
    s.replace_stock(t,pw,stock,"base-orphan-repair",s.state(t)["revision"],source_name="orphan.xlsx")
    with s.engine.begin() as c:
        c.execute(delete(s.tables["baseline_snapshots"]))
        assert c.execute(select(func.count()).select_from(s.tables["stock_state"])).scalar_one()==1
    s.initialize(password=pw)
    with s.engine.connect() as c:
        assert c.execute(select(func.count()).select_from(s.tables["stock_state"])).scalar_one()==0
        assert c.execute(select(func.count()).select_from(s.tables["audit_log"]).where(s.tables["audit_log"].c.action=="REPAIR_ORPHAN_STOCK_WITHOUT_BASELINE")).scalar_one()==1

def test_invoice_ui_hides_draft_workflow_and_cleans_failed_scans():
    source=(Path(__file__).resolve().parents[1]/"inventory_tracker.py").read_text()
    assert 'section("Drafts")' not in source
    assert 'section("Invoice review")' in source
    assert 't("Save changes")' in source
    assert 'failed_empty=[d for d in drafts' in source
    assert 'store.discard_draft(token,draft_id)' in source


def test_scan_lock_is_shared():
    assert _ocr_scan_lock() is _ocr_scan_lock()


def test_worker_invalid_upload():
    with pytest.raises(ValueError):_run_ocr_worker(b"")


def test_worker_missing_file(tmp_path):
    with pytest.raises(RuntimeError):_run_ocr_worker(b"image",worker_path=tmp_path/"missing.py")


def test_worker_success_without_model(tmp_path,monkeypatch):
    import gp_ocr
    monkeypatch.setattr(gp_ocr,"_ocr_memory_state",lambda:(None,None))
    worker=tmp_path/"worker.py"
    worker.write_text('import sys,json\nfrom pathlib import Path\nPath(sys.argv[2]).write_text(json.dumps({"ok":True,"data":{"items":[]}}))')
    assert _run_ocr_worker(b"image",worker_path=worker,timeout=3)=={"items":[]}


def test_worker_timeout(tmp_path,monkeypatch):
    import gp_ocr
    monkeypatch.setattr(gp_ocr,"_ocr_memory_state",lambda:(None,None))
    worker=tmp_path/"worker.py";worker.write_text('import time\ntime.sleep(30)')
    with pytest.raises(RuntimeError,match="timed out"):_run_ocr_worker(b"image",worker_path=worker,timeout=.15)


def test_ui_uses_sidebar_without_global_rtl_rule():
    root=Path(__file__).resolve().parents[1]
    source=(root/"inventory_tracker.py").read_text()
    assert "with st.sidebar:" in source and "sidebar_navigation" in source
    assert 'requested=st.session_state.pop("nav_request",None)' in source
    assert 'if "sidebar_nav" not in st.session_state or requested in nav:' in source
    assert 'st.session_state.get("sidebar_nav")!=desired' not in source
    assert 'st.session_state["nav_request"]="Invoices"' in source
    assert 'st.session_state["nav_request"]="Movements"' in source
    assert 'st.session_state["nav_request"]="Stock"' in source
    assert 'st.session_state["nav_request"]="Analysis"' in source
    assert "use_container_width" not in source
    stylesheet=(root/"assets/style.css").read_text()
    assert "h1, h2, h3, h4, p, span, label, div" not in stylesheet
    assert '[data-testid="stSidebar"]' in stylesheet
    for f in root.glob("*.py"):ast.parse(f.read_text())


def test_web_process_has_no_neural_network_imports():
    root=Path(__file__).resolve().parents[1]
    for file in ("inventory_tracker.py","gp_ocr.py","gp_store.py","gp_core.py"):
        tree=ast.parse((root/file).read_text())
        imports=[]
        for node in ast.walk(tree):
            if isinstance(node,ast.Import):imports += [n.name for n in node.names]
            if isinstance(node,ast.ImportFrom):imports.append(node.module or "")
        assert not any(x.split(".")[0] in ("torch","easyocr","torchvision","rapidocr","onnxruntime") for x in imports)


def test_requirements_remove_torch_and_easyocr():
    text=(Path(__file__).resolve().parents[1]/"requirements.txt").read_text().lower()
    assert "easyocr" not in text and "torch @" not in text and "torchvision" not in text
    assert "rapidocr" in text and "onnxruntime" in text


def test_daily_report_preserves_sections_and_totals():
    from gp_reports import day_report_sheets
    from decimal import Decimal
    ledger=pd.DataFrame([
        {'ledger_id':1,'item_key':'A','item_code':'010716','item_name':'Camera',
         'movement_type':'OUT','quantity':Decimal('3'),'quantity_before':Decimal('10'),
         'quantity_after':Decimal('7'),'without_invoice':True},
        {'ledger_id':2,'item_key':'A','item_code':'010716','item_name':'Camera',
         'movement_type':'IN','quantity':Decimal('1'),'quantity_before':Decimal('7'),
         'quantity_after':Decimal('8'),'without_invoice':False}])
    stock=ensure_unique_stock_keys(pd.DataFrame([{COL_CODE:'010716',COL_NAME:'Camera',COL_QTY:8}]))
    sheets=day_report_sheets('2026-09-08',ledger,stock,pd.DataFrame(),True)
    assert len(sheets)==7
    assert sheets['Daily_Summary'].iloc[0]['total_out']==3
    assert sheets['Daily_Summary'].iloc[0]['total_in']==1
    assert len(sheets['No_Invoice_Alerts'])==1
    row=sheets['Affected_Items'].iloc[0]
    assert row['opening']==10 and row['closing']==8
    assert len(excel_bytes(sheets))>1000


def test_daily_report_empty_is_valid():
    from gp_reports import day_report_sheets
    sheets=day_report_sheets('2026-09-08',pd.DataFrame(),pd.DataFrame(),pd.DataFrame(),False)
    assert sheets['Daily_Summary'].iloc[0]['movements']==0
    assert len(excel_bytes(sheets))>1000




def test_delivery_word_fallback_detects_out_movement():
    boxes=[{"text":"يرجى تسليم المواد المذكورة أدناه","score":.90,"x":500,"y":300,"region":"header"}]
    movement,customer=parse_header_metadata(boxes,1000,1200)
    assert movement=="OUT"


def test_auto_invoice_fields_are_prefilled_but_correctable():
    source=(Path(__file__).resolve().parents[1]/"inventory_tracker.py").read_text()
    assert 'reference=a.text_input(t("Reference"),value=detected_reference,key="reference_"+suffix)' in source
    assert 'kind=d.selectbox(t("Movement type"),kinds,index=kinds.index(current_kind)' in source
    assert 'kind_auto_' not in source
    assert 'if kind not in ("IN","OUT"):raise AppError("Movement type required")' in source
    assert 'duplicate_action=' not in source
    assert 'store.replace_invoice(' in source


def test_invoice_entry_is_only_auto_or_manual():
    source=(Path(__file__).resolve().parents[1]/"inventory_tracker.py").read_text()
    assert 'st.radio(t("Invoice entry"),["AUTO","MANUAL"]' in source
    assert 'key="draft_selector"' not in source
    assert 't("Start manual invoice")' in source


def test_invoice_review_and_closing_stock_compare_present():
    source=(Path(__file__).resolve().parents[1]/"inventory_tracker.py").read_text()
    assert 'reference=a.text_input(t("Reference"),value=detected_reference,key="reference_"+suffix)' in source
    assert 'Invoice number manual fallback' in source
    assert 'if not reference.strip():raise AppError("Invoice number required")' in source
    assert 'Upload closing stock report' in source
    assert 'comparison["Difference"]=comparison["Counted quantity"]-comparison["System quantity"]' in source


def test_mobile_camera_requests_1080p_full_width_and_items_are_editable():
    source=(Path(__file__).resolve().parents[1]/"inventory_tracker.py").read_text()
    assert 'resolution="1080p",width="stretch"' in source
    assert 'st.data_editor(frame,hide_index=True,num_rows="dynamic"' in source
    assert 'disabled=["item_name"]' in source


def test_admin_can_create_user_and_duplicate_is_safe_error():
    s=Store.for_tests();pw="Test-Password-2026";s.initialize(password=pw);t=s.login("admin",pw)
    s.create_user(t,"store1","Another-Test-Password-2026","store","Store User")
    users={u["username"]:u for u in s.list_users(t)}
    assert users["store1"]["role"]=="store"
    assert users["store1"]["display_name"]=="Store User"
    try:
        s.create_user(t,"store1","Another-Test-Password-2026","store","Store User")
    except AppError as error:
        assert str(error)=="Username already exists"
    else:
        raise AssertionError("duplicate username must fail")

def test_safe_error_message_exposes_only_exception_class_and_sqlstate():
    source=(Path(__file__).resolve().parents[1]/"inventory_tracker.py").read_text()
    assert 'safe_detail=kind+(" / SQLSTATE "+sqlstate if sqlstate else "")' in source
    assert 'repr(error)' not in source


def test_new_user_form_has_no_approval_password_and_backend_uses_session_only():
    app=(Path(__file__).resolve().parents[1]/"inventory_tracker.py").read_text()
    store=(Path(__file__).resolve().parents[1]/"gp_store.py").read_text()
    assert 'key="password_useradmin"' not in app
    assert 'store.create_user(token,username,newpass,role,display)' in app
    assert 'def create_user(self,token,username,new_password,role,display_name=""):' in store
    assert 'actor=self._actor(c,token,admin=True)' in store


def test_store_cache_is_versioned_by_build_and_stale_apperror_is_shown_safely():
    source=(Path(__file__).resolve().parents[1]/"inventory_tracker.py").read_text()
    assert 'def get_store(settings_json,build):' in source
    assert 'get_store(json.dumps(config,sort_keys=True),BUILD)' in source
    assert 'type(error).__name__=="AppError"' in source


def test_delete_confirmation_does_not_disable_form_submit_buttons():
    source=(Path(__file__).resolve().parents[1]/"inventory_tracker.py").read_text()
    assert 'Delete invoice"),type="primary",disabled=not confirmed' not in source
    assert 'Delete movement"),type="primary",disabled=not confirmed' not in source
    assert 'Delete movement history report"),type="primary",disabled=not confirmed' not in source
    assert source.count('if not confirmed:raise AppError("Confirm delete")') >= 3


def test_stock_upload_after_daily_movements_switches_to_compare_only():
    source=(Path(__file__).resolve().parents[1]/"inventory_tracker.py").read_text()
    assert 'today_moves=store.ledger(token,store.today())' in source
    assert 'if not today_moves.empty:' in source
    assert 'Stock upload compare mode' in source
    assert '_stock_reconciliation(stock,df,store.today())' in source
    store_source=(Path(__file__).resolve().parents[1]/"gp_store.py").read_text()
    assert 'Set a new baseline before the first movement of the day' in store_source

def test_closing_and_settings_share_stock_reconciliation_helper():
    source=(Path(__file__).resolve().parents[1]/"inventory_tracker.py").read_text()
    assert 'def _stock_reconciliation(expected_stock,counted,day):' in source
    assert '_stock_reconciliation(stock,counted,day)' in source


def test_rapidocr_pin_is_refreshed_for_model_permissions():
    req=(Path(__file__).resolve().parents[1]/"requirements.txt").read_text()
    assert "rapidocr==3.8.4" in req


def test_rapidocr_uses_writable_model_root():
    source=(Path(__file__).resolve().parents[1]/"invoice_ocr_worker.py").read_text()
    assert 'GP_RAPIDOCR_MODEL_ROOT' in source
    assert '/tmp/gp_rapidocr_models' in source
    assert source.count('"Global.model_root_dir": _rapidocr_model_root()')==2
