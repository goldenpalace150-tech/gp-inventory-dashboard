"""On-demand report exports; never rebuilt for every app rerun."""
from datetime import datetime
from decimal import Decimal
import io
import json
import pandas as pd
from gp_core import COL_KEY, COL_MATCH


def visible_frame(frame):
    return frame.drop(columns=[COL_KEY,COL_MATCH],errors="ignore").copy()


def excel_bytes(sheets):
    from openpyxl.styles import Font, PatternFill, Alignment
    result=io.BytesIO()
    with pd.ExcelWriter(result,engine="openpyxl") as writer:
        for title,data in sheets.items():
            df=data.copy() if isinstance(data,pd.DataFrame) else pd.DataFrame(data)
            def safe(v):
                if isinstance(v,(dict,list)):v=json.dumps(v,ensure_ascii=False,default=str)
                if isinstance(v,str) and v[:1] in ("=","+","-","@","\t","\r"):return "'"+v
                if isinstance(v,Decimal):return float(v)
                if isinstance(v,datetime) and v.tzinfo is not None:return v.isoformat()
                return v
            for col in df.columns:df[col]=df[col].map(safe)
            df.to_excel(writer,index=False,sheet_name=title[:31])
            ws=writer.sheets[title[:31]];ws.freeze_panes="A2";ws.sheet_view.rightToLeft=True
            ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="12263E")
                cell.alignment=Alignment(horizontal="right",vertical="center",wrap_text=True)
            ws.row_dimensions[1].height=32
            for cells in ws.columns:
                name=cells[0].column_letter
                maximum=max((len(str(cell.value or "")) for cell in cells[:100]),default=12)
                ws.column_dimensions[name].width=min(46,max(14,maximum+3))
    return result.getvalue()


def day_report_sheets(day,ledger,stock,analysis,closed):
    """Preserve the original daily export sections, using committed records."""
    total_in=float(ledger.loc[ledger['movement_type']=='IN','quantity'].sum()) if not ledger.empty else 0
    total_out=float(ledger.loc[ledger['movement_type']=='OUT','quantity'].sum()) if not ledger.empty else 0
    alerts=ledger[ledger['without_invoice']].copy() if not ledger.empty else pd.DataFrame()
    affected=[]
    if not ledger.empty:
        for key,group in ledger.sort_values('ledger_id').groupby('item_key',sort=False):
            affected.append({'item_code':group.iloc[0]['item_code'],'item_name':group.iloc[0]['item_name'],
                'opening':group.iloc[0]['quantity_before'],
                'total_in':group.loc[group['movement_type']=='IN','quantity'].sum(),
                'total_out':group.loc[group['movement_type']=='OUT','quantity'].sum(),
                'closing':group.iloc[-1]['quantity_after']})
    reorder_col='\u062d\u0627\u0644\u0629 \u0627\u0644\u0637\u0644\u0628'
    clearance_col='\u0627\u0642\u062a\u0631\u0627\u062d \u0627\u0644\u062a\u0635\u0631\u064a\u0641'
    clean=visible_frame(analysis)
    reorders=clean[clean[reorder_col]=='\u0625\u0639\u0627\u062f\u0629 \u0637\u0644\u0628'] if reorder_col in clean else pd.DataFrame()
    clearance=clean[clean[clearance_col]=='\u0645\u0631\u0634\u062d \u0644\u0644\u062a\u0635\u0631\u064a\u0641'] if clearance_col in clean else pd.DataFrame()
    return {'Daily_Summary':pd.DataFrame([{'date':str(day),'closed':bool(closed),'movements':len(ledger),'total_in':total_in,'total_out':total_out,'no_invoice':len(alerts)}]),
            'Daily_Movements':ledger,'Affected_Items':pd.DataFrame(affected),'Closing_Stock':visible_frame(stock),
            'No_Invoice_Alerts':alerts,'Reorder_List':reorders,'Slow_Clearance':clearance}
